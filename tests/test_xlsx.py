import io
import sqlite3
import urllib.parse

import pytest
from datasette.app import Datasette
from openpyxl import load_workbook

import datasette_xlsx


@pytest.fixture
def ds(tmp_path):
    path = str(tmp_path / "test.db")
    conn = sqlite3.connect(path)
    conn.execute("create table personer (id integer primary key, namn text, nr integer, f text)")
    conn.executemany(
        "insert into personer (namn, nr, f) values (?, ?, ?)",
        [(f"Person {i}", 199001011234 + i, "=1+1" if i == 0 else "ö\x01k")
         for i in range(2500)],
    )
    conn.execute("create view v as select * from personer where id % 2 = 0")
    conn.commit()
    return Datasette([path], settings={"max_returned_rows": 1000},
                     metadata={"databases": {"test": {"queries": {
                         "udda": "select id, namn from personer where id % 2 = 1"}}}})


def sheet(resp):
    assert resp.status_code == 200, resp.text
    wb = load_workbook(io.BytesIO(resp.content))
    return wb, wb.worksheets[0]


@pytest.mark.asyncio
async def test_table_all_pages(ds):
    wb, ws = sheet(await ds.client.get("/test/personer.xlsx"))
    assert ws.max_row == 2501
    assert [c.value for c in ws[1]] == ["id", "namn", "nr", "f"]
    assert ws["D2"].value == "=1+1" and ws["D2"].data_type == "s"
    assert ws["D3"].value == "ök"
    assert ws["C2"].value == 199001011234
    assert "Om uttaget" in wb.sheetnames


@pytest.mark.asyncio
async def test_filtered_sorted(ds):
    wb, ws = sheet(await ds.client.get(
        "/test/personer.xlsx?id__gt=2000&_sort_desc=id&_col=namn"))
    assert ws.max_row == 501
    assert [c.value for c in ws[1]] == ["id", "namn"]
    assert ws["A2"].value == 2500


@pytest.mark.asyncio
async def test_view(ds):
    wb, ws = sheet(await ds.client.get("/test/v.xlsx"))
    assert ws.max_row == 1251


@pytest.mark.asyncio
async def test_sql_not_truncated(ds):
    wb, ws = sheet(await ds.client.get(
        "/test.xlsx?sql=select+*+from+personer+where+id+>+:min&min=10"))
    assert ws.max_row == 2491


@pytest.mark.asyncio
async def test_canned(ds):
    wb, ws = sheet(await ds.client.get("/test/udda.xlsx"))
    assert ws.max_row == 1251


@pytest.mark.asyncio
async def test_row(ds):
    wb, ws = sheet(await ds.client.get("/test/personer/5.xlsx"))
    assert ws.max_row == 2 and ws["A2"].value == 5


@pytest.mark.asyncio
async def test_link_on_page(ds):
    html = (await ds.client.get("/test/personer")).text
    assert "personer.xlsx" in html


@pytest.mark.asyncio
async def test_permissions_forwarded(tmp_path):
    path = str(tmp_path / "p.db")
    conn = sqlite3.connect(path)
    conn.execute("create table t (id integer primary key)")
    conn.executemany("insert into t values (?)", [(i,) for i in range(1500)])
    conn.commit()
    ds = Datasette([path], settings={"max_returned_rows": 1000},
                   metadata={"databases": {"p": {"allow": {"id": "henrik"}}}})
    assert (await ds.client.get("/p/t.xlsx")).status_code == 403
    cookies = {"ds_actor": ds.sign({"a": {"id": "henrik"}}, "actor")}
    resp = await ds.client.get("/p/t.xlsx", cookies=cookies)
    wb, ws = sheet(resp)
    assert ws.max_row == 1501


def table_of(ws):
    """Bladets enda Excel-tabell; kontrollerar att rubrikerna matchar."""
    assert len(ws.tables) == 1
    table = list(ws.tables.values())[0]
    header_cells = [c.value for c in ws[1]]
    assert table.column_names == header_cells
    assert table.autoFilter.ref == table.ref
    return table


@pytest.mark.asyncio
async def test_excel_table(ds):
    wb, ws = sheet(await ds.client.get("/test/personer.xlsx"))
    table = table_of(ws)
    assert table.displayName == "tbl_personer"
    assert table.ref == "A1:D2501"
    assert table.tableStyleInfo.name == "TableStyleMedium2"
    assert not wb["Om uttaget"].tables


@pytest.mark.asyncio
async def test_excel_table_invalid_headers(ds):
    sql = 'select id, namn as ID, 1 as "", 2 as "a\nb" from personer limit 3'
    wb, ws = sheet(await ds.client.get(
        "/test.xlsx?" + urllib.parse.urlencode({"sql": sql})))
    table = table_of(ws)
    assert table.column_names == ["id", "ID_2", "Kolumn3", "a b"]
    assert table.displayName == "tbl_test_fraga"
    assert table.ref == "A1:D4"


@pytest.mark.asyncio
async def test_excel_table_empty_result(ds):
    wb, ws = sheet(await ds.client.get("/test/personer.xlsx?id__gt=99999"))
    assert ws.max_row == 1
    assert table_of(ws).ref == "A1:D2"


@pytest.mark.asyncio
async def test_excel_table_per_sheet(ds, monkeypatch):
    monkeypatch.setattr(datasette_xlsx, "EXCEL_MAX_ROWS", 1000)
    wb, ws = sheet(await ds.client.get("/test/personer.xlsx"))
    data_sheets = [s for s in wb.worksheets if s.title != "Om uttaget"]
    tables = [table_of(s) for s in data_sheets]
    assert [t.displayName for t in tables] == [
        "tbl_personer", "tbl_personer_2", "tbl_personer_3"]
    assert [t.ref for t in tables] == ["A1:D1000", "A1:D1000", "A1:D503"]


def widths(ws):
    return {k: d.width for k, d in ws.column_dimensions.items() if d.width}


@pytest.mark.asyncio
async def test_column_widths(ds):
    wb, ws = sheet(await ds.client.get("/test/personer.xlsx"))
    # id: rubrik + filterknapp, namn: "Person 999" (bara de första 1000
    # raderna ingår i beräkningen), nr: Allmänt visar
    # max 11 tecken, f: minimibredd
    assert widths(ws) == {"A": 6, "B": 11, "C": 12, "D": 6}
    # Raderna kommer i rätt ordning även efter bufferten för breddberäkning
    assert [ws.cell(r, 1).value for r in range(1000, 1004)] == [999, 1000, 1001, 1002]
    assert widths(wb["Om uttaget"])["A"] > 6


@pytest.mark.asyncio
async def test_column_widths_capped_and_multiline(ds):
    sql = "select 'x' || char(10) || 'kort' as m, printf('%.200c', 'a') as lang"
    wb, ws = sheet(await ds.client.get(
        "/test.xlsx?" + urllib.parse.urlencode({"sql": sql})))
    assert widths(ws) == {"A": 6, "B": 60}


@pytest.mark.asyncio
async def test_column_widths_every_sheet(ds, monkeypatch):
    monkeypatch.setattr(datasette_xlsx, "EXCEL_MAX_ROWS", 600)
    wb, ws = sheet(await ds.client.get("/test/personer.xlsx"))
    data_sheets = [s for s in wb.worksheets if s.title != "Om uttaget"]
    assert len(data_sheets) == 5
    assert all(widths(s) == widths(ws) for s in data_sheets)
    assert sum(s.max_row - 1 for s in data_sheets) == 2500
