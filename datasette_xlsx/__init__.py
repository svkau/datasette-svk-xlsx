"""
datasette-xlsx – Excel-uttag för Datasette.

Registrerar en output renderer för filändelsen .xlsx. Samma mekanism som
Datasettes inbyggda .json, så länken dyker upp automatiskt på tabell-, vy-,
fråge- och radsidor.

Hela resultatet hämtas, oberoende av sidindelning:

* Tabeller/vyer: pluginen ber Datasette själv om sidorna (.json med
  _size=max) och följer "next"-token. All filtrering (__exact, _where,
  _search, _sort, _col/_nocol, facetter …) görs alltså av Datasette, inte
  av pluginen. Behörigheter följer med eftersom anroparens cookies och
  Authorization-header vidarebefordras.
* SQL-frågor och sparade frågor: pagineras inte i Datasette utan trunkeras
  vid max_returned_rows. Här körs frågan om direkt mot en läsanslutning
  (skrivskyddad) och rader strömmas in i arbetsboken.

Datat i varje blad formateras som en Excel-tabell (filter, randade rader).
Kolumnnamn som Excel inte godtar som tabellrubriker (tomma eller dubbletter)
görs om, t.ex. id, id → id, id_2.

Kolumnbredder beräknas från rubrikerna och de första 1000 raderna (xlsx
saknar automatisk anpassning vid öppning), med tak på 60 tecken.

Konfiguration (metadata/datasette.yaml):

    plugins:
      datasette-xlsx:
        max_rows: 1000000      # tak för antal rader totalt
        time_limit_ms: 30000   # tidsgräns för SQL-frågor
"""
import datetime
import re
import tempfile
import urllib.parse
import warnings

from datasette import hookimpl
from datasette.utils import sqlite_timelimit
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

PLUGIN = "datasette-xlsx"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXCEL_MAX_ROWS = 1_048_576  # inkl. rubrikrad
EXCEL_MAX_CELL = 32_767
EXCEL_MAX_INT = 10**15  # Excel lagrar tal som double, ~15 siffrors precision
TABLE_STYLE = "TableStyleMedium2"
WIDTH_SAMPLE_ROWS = 1000  # rader som kolumnbredden beräknas från
MIN_WIDTH = 6
MAX_WIDTH = 60
FILTER_BUTTON_WIDTH = 3  # plats för tabellens filterknapp i rubriken


@hookimpl
def register_output_renderer(datasette):
    return {"extension": "xlsx", "render": render_xlsx, "can_render": can_render}


def can_render(view_name):
    return view_name in ("table", "database", "query", "row")


class _Writer:
    """Strömmar rader till en write-only-arbetsbok och hanterar Excels gränser."""

    def __init__(self, columns, sheet_title, max_rows):
        self.wb = Workbook(write_only=True)
        self.columns = list(columns)
        self.headers = _table_headers(self.columns)
        self.base_title = _sheet_title(sheet_title)
        self.max_rows = max_rows
        self.rows = 0
        self.sheets = 0
        self.truncated_cells = 0
        self.hit_limit = False
        self.ws = None
        self.table = None
        self.widths = None  # beräknas från de första raderna i första bladet
        self.buffer = None
        self._new_sheet()

    def _new_sheet(self):
        self._flush()
        self._close_table()
        self.sheets += 1
        title = self.base_title if self.sheets == 1 else _sheet_title(
            f"{self.base_title[:25]} ({self.sheets})"
        )
        self.ws = self.wb.create_sheet(title)
        self.ws.freeze_panes = "A2"
        self.sheet_rows = 1
        if self.widths is None:
            # Kolumnbredder måste sättas innan första raden skrivs, så rader
            # buffras tills underlaget räcker (se _flush).
            self.buffer = []
        else:
            self._start_sheet()
        if self.headers:
            self.table = Table(
                displayName=_table_name(self.base_title, self.sheets),
                ref="A1:A2",  # sätts slutligt i _close_table
                autoFilter=AutoFilter(ref="A1:A2"),
                tableStyleInfo=TableStyleInfo(name=TABLE_STYLE, showRowStripes=True),
            )
            # I write-only-läge kan openpyxl inte läsa rubrikerna ur bladet,
            # så kolumnerna anges explicit och måste matcha rubrikcellerna.
            self.table.tableColumns = [
                TableColumn(id=i, name=h) for i, h in enumerate(self.headers, 1)
            ]
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                self.ws.add_table(self.table)

    def _start_sheet(self):
        """Sätt kolumnbredder och skriv rubrikraden i aktuellt blad."""
        _set_widths(self.ws, self.widths)
        self.ws.append([self._cell(h) for h in self.headers])

    def _flush(self):
        """Beräkna kolumnbredder från buffrade rader och skriv ut dem."""
        if self.buffer is None:
            return
        self.widths = _column_widths(self.headers, self.buffer)
        self._start_sheet()
        for row in self.buffer:
            self.ws.append([self._cell(v) for v in row])
        self.buffer = None

    def _close_table(self):
        """Sätt slutligt område för aktuellt blads Excel-tabell."""
        if self.table is None:
            return
        # En tabell måste ha minst en datarad; vid tomt resultat blir den tom.
        last_row = max(self.sheet_rows, 2)
        self.table.ref = f"A1:{get_column_letter(len(self.headers))}{last_row}"
        self.table.autoFilter.ref = self.table.ref
        self.table = None

    def _cell(self, value):
        if value is None:
            return None
        if isinstance(value, bool):
            return value
        if isinstance(value, int):
            if abs(value) >= EXCEL_MAX_INT:
                value = str(value)  # bevara exakta siffror
            else:
                return value
        elif isinstance(value, float):
            return value
        elif isinstance(value, (bytes, bytearray, memoryview)):
            value = f"<binärdata, {len(value)} byte>"
        elif isinstance(value, dict) and "value" in value:
            # _labels=on: {"value": ..., "label": ...}
            value = value.get("label") or value.get("value")
            return self._cell(value)
        elif not isinstance(value, str):
            value = str(value)
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > EXCEL_MAX_CELL:
            value = value[: EXCEL_MAX_CELL - 20] + " …[TRUNKERAT]"
            self.truncated_cells += 1
        cell = WriteOnlyCell(self.ws, value=value)
        # Tvinga text så att "=..." aldrig tolkas som formel
        cell.data_type = "s"
        return cell

    def write(self, row):
        """Returnerar False när max_rows nåtts."""
        if self.rows >= self.max_rows:
            self.hit_limit = True
            return False
        if self.sheet_rows >= EXCEL_MAX_ROWS:
            self._new_sheet()
        if self.buffer is not None:
            self.buffer.append(list(row))
            if len(self.buffer) >= WIDTH_SAMPLE_ROWS:
                self._flush()
        else:
            self.ws.append([self._cell(v) for v in row])
        self.rows += 1
        self.sheet_rows += 1
        return True

    def finish(self, info):
        self._flush()
        self._close_table()
        ws = self.wb.create_sheet("Om uttaget")
        info = dict(info)
        info["Antal rader"] = self.rows
        info["Antal blad med data"] = self.sheets
        if self.hit_limit:
            info["OBS"] = f"Uttaget avbröts vid max_rows = {self.max_rows}"
        if self.truncated_cells:
            info["Trunkerade celler"] = (
                f"{self.truncated_cells} (Excel tillåter max {EXCEL_MAX_CELL} tecken per cell)"
            )
        info = [(k, v) for k, v in info.items() if v not in (None, "")]
        _set_widths(ws, _column_widths(["", ""], info, max_width=100))
        for k, v in info:
            ws.append([k, self._cell(v)])
        with tempfile.TemporaryFile() as fp:
            self.wb.save(fp)
            fp.seek(0)
            return fp.read()


def _sheet_title(name):
    name = re.sub(r"[\[\]:*?/\\]", "_", name or "Resultat").strip("'") or "Resultat"
    return name[:31]


def _text_width(value):
    """Ungefärlig visningsbredd i tecken för ett cellvärde."""
    if value is None:
        return 0
    if isinstance(value, bool):
        return 5
    if isinstance(value, (int, float)):
        # Formatet Allmänt visar högst 11 tecken innan Excel byter till
        # exponentform; större heltal skrivs som text av _cell.
        text = str(value)
        return len(text) if abs(value) >= EXCEL_MAX_INT else min(len(text), 11)
    if isinstance(value, (bytes, bytearray, memoryview)):
        return len(f"<binärdata, {len(value)} byte>")
    if isinstance(value, dict) and "value" in value:
        return _text_width(value.get("label") or value.get("value"))
    return max(len(line) for line in str(value).splitlines() or [""])


def _column_widths(headers, rows, max_width=MAX_WIDTH):
    """Kolumnbredder utifrån rubriker och ett urval rader."""
    widths = [len(h) + FILTER_BUTTON_WIDTH if h else 0 for h in headers]
    for row in rows:
        for i, value in enumerate(row[: len(widths)]):
            widths[i] = max(widths[i], _text_width(value))
    return [min(max(w + 1, MIN_WIDTH), max_width) for w in widths]


def _set_widths(ws, widths):
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width


def _table_headers(columns):
    """Gör kolumnnamn giltiga som rubriker i en Excel-tabell.

    Excel kräver icke-tomma rubriker som är unika utan hänsyn till
    versaler/gemener, t.ex. blir ``select a.id, b.id`` rubrikerna id, id_2.
    """
    headers, seen = [], set()
    for i, col in enumerate(columns, 1):
        name = ILLEGAL_CHARACTERS_RE.sub("", str(col))
        name = re.sub(r"\s+", " ", name).strip()[:255] or f"Kolumn{i}"
        candidate, n = name, 1
        while candidate.casefold() in seen:
            n += 1
            candidate = f"{name}_{n}"
        seen.add(candidate.casefold())
        headers.append(candidate)
    return headers


def _table_name(base, sheet_number):
    """Giltigt och unikt namn för Excel-tabellen, t.ex. tbl_personer_2."""
    name = "tbl_" + re.sub(r"[^\w.]", "_", base)[:200]
    return name if sheet_number == 1 else f"{name}_{sheet_number}"


def _filename_header(name):
    name = re.sub(r'[\\/:*?"<>|\r\n]+', "_", name) + ".xlsx"
    ascii_name = name.encode("ascii", "replace").decode().replace("?", "_")
    quoted = urllib.parse.quote(name)
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quoted}"


def _forward_auth(request):
    headers = {}
    auth = request.headers.get("authorization")
    if auth:
        headers["authorization"] = auth
    return {"cookies": request.cookies, "headers": headers}


async def _table_rows(datasette, request, columns, writer):
    """Följ Datasettes egen keyset-paginering via interna .json-anrop."""
    path = request.path[: -len(".xlsx")] + ".json"
    args = [
        (k, v)
        for k, v in urllib.parse.parse_qsl(request.query_string, keep_blank_values=True)
        if k not in ("_next", "_size", "_shape", "_format", "_extra")
    ]
    args += [("_size", "max"), ("_shape", "objects")]
    auth = _forward_auth(request)
    next_token = None
    while True:
        qs = args + ([("_next", next_token)] if next_token else [])
        response = await datasette.client.get(
            path + "?" + urllib.parse.urlencode(qs), **auth
        )
        if response.status_code != 200:
            raise RuntimeError(f"{response.status_code} från {path}: {response.text[:200]}")
        data = response.json()
        for row in data.get("rows", []):
            if not writer.write([row.get(c) for c in columns]):
                return
        next_token = data.get("next")
        if not next_token:
            return


async def _query_rows(db, sql, params, writer, time_limit_ms):
    """Kör frågan på nytt utan max_returned_rows och strömma raderna."""

    def run(conn):
        with sqlite_timelimit(conn, time_limit_ms):
            cursor = conn.execute(sql, params or {})
            for row in cursor:
                if not writer.write(row):
                    break

    await db.execute_fn(run)


async def render_xlsx(datasette, request, database, table, sql, query_name,
                      columns, rows, view_name, data):
    config = datasette.plugin_config(PLUGIN, database=database, table=table) or {}
    max_rows = int(config.get("max_rows", 1_000_000))
    time_limit_ms = int(config.get("time_limit_ms", 30_000))
    db = datasette.get_database(database)
    name = table or query_name or f"{database}-fraga"
    writer = _Writer(columns, name, max_rows)

    if view_name == "row":
        for row in rows:
            writer.write([row[c] for c in columns] if isinstance(row, dict) else row)
    elif table and view_name == "table":
        await _table_rows(datasette, request, columns, writer)
    else:
        if query_name:
            canned = await datasette.get_canned_query(
                database, query_name, request.actor
            )
            if canned and canned.get("write"):
                return {"body": "Skrivande frågor kan inte exporteras",
                        "status_code": 400}
        params = (data.get("query") or {}).get("params") or {}
        await _query_rows(db, sql, params, writer, time_limit_ms)

    info = {
        "Databas": database,
        "Tabell/vy": table,
        "Sparad fråga": query_name,
        "SQL": sql if view_name != "table" else None,
        "URL": request.url.replace(".xlsx", "", 1),
        "Användare": (request.actor or {}).get("id"),
        "Exporterad": datetime.datetime.now().astimezone().isoformat(timespec="seconds"),
        "Datasette-version": __import__("datasette").__version__,
    }
    body = writer.finish(info)
    return {
        "body": body,
        "content_type": CONTENT_TYPE,
        "headers": {"content-disposition": _filename_header(name)},
    }
